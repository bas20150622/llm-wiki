import pytest

from llm_wiki.convert import convert, frontmatter_block, read_frontmatter, sha256


def test_frontmatter_block_round_trips(tmp_path):
    md = tmp_path / "r.md"
    md.write_text(frontmatter_block("r.pdf", "abc") + "body\n")
    meta = read_frontmatter(md)
    assert meta["type"] == "converted-source"
    assert meta["converted_from"] == "r.pdf"
    assert meta["source_sha256"] == "abc"
    assert meta["converted_by"].startswith("docling/")


def test_skips_when_hash_matches(tmp_path, capsys):
    pdf = tmp_path / "r.pdf"
    pdf.write_bytes(b"%PDF-1.4 test")
    (tmp_path / "r.md").write_text(frontmatter_block("r.pdf", sha256(pdf)) + "body\n")
    convert(str(pdf))
    assert "Up to date" in capsys.readouterr().out


@pytest.mark.parametrize(
    "existing",
    [
        "hand written\n",                                    # no frontmatter
        frontmatter_block("r.docx", "x") + "body\n",         # stem collision
        frontmatter_block("r.pdf", "stale-hash") + "body\n",  # native file changed
    ],
)
def test_refuses_to_overwrite(tmp_path, existing):
    pdf = tmp_path / "r.pdf"
    pdf.write_bytes(b"%PDF-1.4 test")
    (tmp_path / "r.md").write_text(existing)
    with pytest.raises(SystemExit) as exit_info:
        convert(str(pdf))
    assert exit_info.value.code == 2
    assert (tmp_path / "r.md").read_text() == existing


def test_rejects_unsupported_and_missing(tmp_path):
    with pytest.raises(SystemExit):
        convert(str(tmp_path / "x.txt"))
    with pytest.raises(SystemExit):
        convert(str(tmp_path / "missing.pdf"))


@pytest.mark.slow
def test_real_pptx_conversion_has_slide_markers(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    for number in (1, 2):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Slide {number}"
    path = tmp_path / "deck.pptx"
    deck.save(path)

    convert(str(path))
    text = (tmp_path / "deck.md").read_text()
    assert text.startswith("---\ntype: \"converted-source\"")
    assert "<!-- slide 1 -->" in text and "<!-- slide 2 -->" in text
